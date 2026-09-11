"""
Build an editable Word (.docx) version of the HFN housemanship policy brief.

Output: docs/hfn-housemanship-policy-brief.docx

Deliberately clean and lightly styled so HFN leadership can edit freely and drop
it onto HFN letterhead. Content mirrors docs/hfn-housemanship-policy-brief.md.

Run:
  python3 scripts/build-hfn-housemanship-docx.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "hfn-housemanship-policy-brief.docx"

GREEN = RGBColor(0x1B, 0x5E, 0x3A)
INK = RGBColor(0x14, 0x26, 0x1C)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def build():
    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for lvl, (size, color, before) in {
        "Heading 1": (15, GREEN, 14),
        "Heading 2": (12, INK, 9),
    }.items():
        st = doc.styles[lvl]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)

    def para(text=None, *, bold=False, italic=False, size=None, color=None,
             align=None, space_after=None, space_before=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if space_before is not None:
            p.paragraph_format.space_before = Pt(space_before)
        if text:
            add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
        return p

    def add_runs(p, text, bold=False, italic=False, size=None, color=None):
        # simple **bold** parser
        parts = text.split("**")
        for i, chunk in enumerate(parts):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        add_runs(p, text)
        return p

    def hrule():
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "B8863B")
        pbdr.append(bottom)
        pPr.append(pbdr)
        p.paragraph_format.space_after = Pt(4)

    def shaded_block(lines):
        """A single-cell shaded table used for the address / callout blocks."""
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.paragraphs[0].text = ""
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F2F5F3")
        cell._tc.get_or_add_tcPr().append(shade)
        first = True
        for line in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, line)
            first = False
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    # ---------------- TITLE ----------------
    eyebrow = para("HEALTHCARE FEDERATION OF NIGERIA", bold=True, size=10.5, color=GREEN)
    eyebrow.paragraph_format.space_after = Pt(2)
    note = para("[ Draft for adoption. Place on HFN letterhead before transmission. ]",
                italic=True, size=9, color=MUTED)
    note.paragraph_format.space_after = Pt(10)

    title = para("Unlocking the House Officer Bottleneck", bold=True, size=22, color=INK)
    title.paragraph_format.space_after = Pt(2)
    para("A policy brief and set of recommendations from the Healthcare Federation of Nigeria",
         bold=True, size=12.5, color=GREEN, space_after=10)
    hrule()

    shaded_block([
        "**For the attention of:**  Professor Iziaq Adekunle Salako, Honourable Minister of "
        "State for Health and Social Welfare, Federal Republic of Nigeria",
        "**Copied to:**  The Coordinating Minister of Health and Social Welfare; the Registrar, "
        "Medical and Dental Council of Nigeria",
        "**Prepared for the Healthcare Federation of Nigeria by:**  Dr Debo Odulana, member, "
        "Healthcare Federation of Nigeria",
        "**Status:**  Submitted to HFN leadership for review and adoption",
        "**Date:**  July 2026",
    ])

    # ---------------- EXEC SUMMARY ----------------
    doc.add_heading("Executive summary", level=1)
    para("Nigeria is training more doctors than at any point in its history, and at the same "
         "time turning thousands of them away at the very first step of their careers. Roughly "
         "6,000 doctors qualify each year, but the centralised housemanship system has room for "
         "only about 4,000. Around 2,000 fully qualified graduates are therefore left stranded "
         "every year, unable to begin the one year of supervised internship that stands between "
         "graduation and full registration. Many wait six to twelve months. Some are asked for "
         "bribes and other improper fees to secure a place.")
    para("The bribery and the frustration with the recent centralisation of placement are real, "
         "but they are symptoms. The underlying problem is structural, and it has four distinct "
         "parts: the pool of accredited training capacity is too small and too narrowly defined; "
         "graduate numbers are rising by deliberate government policy against that frozen pool; "
         "the funding of the intern post is fragile and does not follow the graduate; and the "
         "allocation mechanism is centralised, opaque, and easy to game. Left unaddressed, the "
         "gap will widen sharply, because the expansion of medical school intake now underway "
         "will push annual output well above 6,000 while internship capacity stays where it is.")
    para("The Medical and Dental Council of Nigeria has itself told the National Assembly that "
         "the answer is to bring state and private hospitals into the system. The Healthcare "
         "Federation of Nigeria agrees, and goes further. This brief sets out four connected "
         "reforms, each targeted at one of the four root causes, that together would clear the "
         "backlog, protect training quality, and turn a bottleneck into an on-ramp for the "
         "workforce Nigeria is investing so heavily to create.")

    # ---------------- 1 ----------------
    doc.add_heading("1.  The problem, stated plainly", level=1)
    para("A Nigerian medical graduate cannot practise, cannot be fully registered, and in most "
         "cases cannot proceed to residency or leave for national service until completing a "
         "twelve-month housemanship: three months each, uninterrupted, in Internal Medicine, "
         "Surgery, Obstetrics and Gynaecology, and Paediatrics, under consultant supervision in "
         "an accredited hospital. That single requirement has become the tightest chokepoint in "
         "the entire medical workforce pipeline.")
    bullet("About **6,000 doctors qualify each year**; the centralised housemanship system has "
           "capacity for about **4,000**. Roughly **2,000 graduates are stranded annually**. "
           "These are the Council's own figures, presented by the MDCN Registrar to the Senate "
           "Committee on Health.")
    bullet("The stranded graduates are not failing candidates. They have passed every "
           "examination and met every academic requirement. They are simply told there is no "
           "room.")
    bullet("Waiting times of **six to twelve months** are now common, and competition for the "
           "limited slots has created space for **unofficial payments and other improper fees**. "
           "Graduates have publicly reported being asked for sums such as **N400,000** to secure "
           "a place, and describe sought-after centres where about 30 slots draw over 300 "
           "applicants. These pressures fall hardest on graduates from families least able to pay.")
    para("The human cost is a cohort of young, fully qualified doctors kept idle at the exact "
         "moment the country needs them most, and at the exact moment they are deciding whether "
         "their future is in Nigeria or abroad.")

    # ---------------- 2 ----------------
    doc.add_heading("2.  Why this matters now more than ever", level=1)
    para("**First, the pipeline is being deliberately widened at the top.** As part of the "
         "Government's own workforce strategy, admission quotas for medical training have risen "
         "by around 38 per cent since 2023, the number of medical schools has grown, and roughly "
         "twenty more are in the accreditation pipeline. New intakes are reported to have risen "
         "by about 160 per cent between 2023 and the end of 2025. This is the right ambition. But "
         "every one of those additional students becomes a graduate who needs a housemanship "
         "place. If internship capacity stays frozen at 4,000 while output climbs toward 8,000 "
         "and beyond, the number stranded each year does not creep up, it multiplies.")
    para("**Second, the bottleneck feeds the very emigration the Government is fighting.** More "
         "than 8,500 Nigerian doctors joined the United Kingdom register between 2021 and 2024 "
         "alone, and by some measures around half of Nigerian-trained doctors now practise "
         "abroad. A graduate forced to sit idle for a year, hustling for a placement and "
         "sometimes asked to pay for it, receives a powerful early signal that the system does "
         "not value them. Many act on it. Each departure also writes off a public training "
         "investment estimated in the tens of thousands of dollars per doctor. The housemanship "
         "bottleneck is not separate from the japa crisis; it is one of its first and most "
         "avoidable causes.")
    shaded_block([
        "**Clearing the housemanship bottleneck does not require training a single additional "
        "doctor. It requires letting the doctors we have already trained start work. Few reforms "
        "in the health sector offer so much return for so little cost.**",
    ])

    # ---------------- 3 ----------------
    doc.add_heading("3.  Root-cause diagnosis", level=1)
    para("The symptoms, bribery, delay, and the friction of centralised posting, all trace back "
         "to four distinct and addressable causes. They are set out separately below because "
         "each needs its own remedy, and together they account for the whole problem.")

    doc.add_heading("Cause 1:  Accredited capacity is too small and too narrowly defined", level=2)
    para("The training pool is dominated by federal institutions. Of the 114 hospitals "
         "accredited for housemanship, 47 are teaching hospitals, 21 are Federal Medical Centres, "
         "37 are general and specialist hospitals that are mostly state-owned and military, and "
         "just 9 are private. Only 44, all of them federal, sit on the centralised posting "
         "portal; the other 70 accredited hospitals, state, military, and private, are left "
         "outside it entirely. Between them, those 44 federal centres advertise only about 2,783 "
         "medical housemanship places in a posting cycle. The private sector, which now delivers "
         "a large and rising share of Nigerian healthcare, is barely represented, at 9 of 114 "
         "accredited centres.")
    para("Two design choices keep the private and state pool small:")
    bullet("**The accreditation criteria for private hospitals are disproportionately "
           "restrictive.** Requirements around infrastructure, resident accommodation, and "
           "departmental scope are set at a level that many capable, high-volume private "
           "hospitals cannot meet, even where the clinical training they could offer is "
           "excellent. The bar screens out good trainers rather than bad ones.")
    bullet("**The slots allocated to those private hospitals that do qualify are too few.** Even "
           "an accredited private hospital is typically granted a small quota, far below the "
           "number of house officers it could competently supervise. The pool is both shallow "
           "and under-used.")
    para("A related rule compounds the problem: accreditation generally expects a single "
         "hospital to provide all four core postings under one roof. Many otherwise strong "
         "hospitals have real depth in two or three of the four disciplines but not all four, and "
         "are therefore excluded entirely.")

    doc.add_heading("Cause 2:  Graduate demand is rising by policy while capacity stays fixed", level=2)
    para("The Government is expanding medical school intake as a matter of deliberate strategy. "
         "There is currently no mechanism that requires internship capacity to expand in step "
         "with graduate output. Capacity planning for training posts is not tied to the number "
         "of doctors the system is about to produce. The two numbers are managed by different "
         "processes and are drifting apart.")

    doc.add_heading("Cause 3:  The funding of the intern post is fragile and non-portable", level=2)
    para("House officers are paid on the first grade of the Consolidated Medical Salary "
         "Structure, funded through the federal payroll for federal centres. In practice this "
         "funding is unreliable and non-portable:")
    bullet("Intern salary lines are frequently not captured cleanly in institutional annual "
           "budgets, so hospitals decline to take house officers even where the clinical need is "
           "obvious.")
    bullet("Where posts are funded, payment is often delayed; house officers have publicly "
           "reported months of unpaid salaries.")
    bullet("Because the money is attached to specific federal institutions rather than to the "
           "graduate, a willing state or private hospital that could train an intern has no "
           "funding stream to do so.")

    doc.add_heading("Cause 4:  Allocation is centralised, opaque, and easy to game", level=2)
    para("The centralised posting system was introduced to bring order and fairness to "
         "placement. In practice, applied to a supply that is far too small, it has rationed "
         "scarcity rather than expanded access. Vacancies and quotas are not transparently "
         "published in real time, state hospitals sit largely outside the portal, and graduates "
         "compete blind for a shrinking set of preferred slots. The transparently contestable "
         "pool is narrower still than the headline numbers suggest, because even within the "
         "federal centres a large share of places is reserved for discretionary allocation by "
         "hospital managers. Teaching hospitals commonly route only around 40 per cent of their "
         "places to the open portal, reserving the rest for the Chief Medical Director, the "
         "Provost, and internal use. That discretionary space is precisely where graduates "
         "report slots being sold. Opacity plus scarcity is the condition in which unofficial "
         "payments flourish. Centralising the allocation of an inadequate supply did not create "
         "the shortage, but it concentrated and exposed it, and made rent-seeking easier.")

    # MECE table
    doc.add_heading("The four causes and their remedies, at a glance", level=2)
    mece = [
        ("Root cause", "Remedy"),
        ("1.  Accredited supply too small and too narrow (private bar too high, slots too few, "
         "single-roof rule)",
         "Broaden supply safely: proportionate criteria, rotational consortium accreditation, "
         "larger private and state quotas, national quality standard"),
        ("2.  Graduate demand rising by policy against a frozen pool",
         "Plan internship capacity to the graduate pipeline"),
        ("3.  Intern funding fragile and non-portable",
         "Ring-fence a portable house-officer stipend that follows the graduate"),
        ("4.  Allocation centralised, opaque, gameable (bribes, delay)",
         "Transparent, rules-based matching with anti-corruption safeguards"),
    ]
    table = doc.add_table(rows=len(mece), cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (a, b) in enumerate(mece):
        for j, txt in enumerate((a, b)):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(txt)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "1B5E3A")
                cell._tc.get_or_add_tcPr().append(shade)
            r.font.size = Pt(10)
    doc.add_paragraph()

    # ---------------- 4 ----------------
    doc.add_heading("4.  Recommendations", level=1)
    para("The four reforms below map directly onto the four causes. Each can begin within the "
         "current budget cycle. Together they are designed to be comprehensive: fix all four and "
         "the bottleneck clears; fix only some and the pressure simply moves.")

    doc.add_heading("Recommendation 1:  Broaden accredited supply, safely  (targets Cause 1)", level=2)
    bullet("**Right-size the accreditation criteria.** Move to proportionate, competency- and "
           "outcome-based standards focused on what protects trainees: adequate patient volume "
           "and case mix, consultant supervision, structured teaching, and record-keeping. Relax "
           "requirements that exclude capable hospitals without improving training, and review "
           "the private-hospital criteria specifically against this test.")
    bullet("**Adopt a rotational, consortium model of accreditation.** Where no single hospital "
           "offers all four core disciplines to standard, allow two or three hospitals with "
           "complementary strengths to be accredited together as a training network, with house "
           "officers rotating between them under a lead coordinating institution. Housemanship "
           "does not have to happen under one roof.")
    bullet("**Raise the slot quotas** of accredited private and state hospitals to match their "
           "genuine supervisory capacity, and bring all accredited state hospitals onto the "
           "portal.")
    bullet("**Guarantee quality with a national standard.** Introduce a national housemanship "
           "curriculum, a structured logbook, and an end-of-internship competency assessment, so "
           "a house officer trained in a private or state hospital, or across a consortium, meets "
           "the same standard as one trained in a federal teaching hospital. This standard is "
           "what makes expansion defensible.")

    doc.add_heading("Recommendation 2:  Plan internship capacity to the graduate pipeline  "
                    "(targets Cause 2)", level=2)
    bullet("Establish a simple, published, rolling reconciliation between projected annual "
           "graduate output and accredited housemanship capacity, reviewed each year by MDCN and "
           "the Ministry together.")
    bullet("Make the approval of new medical schools and intake increases conditional on a "
           "credible plan for the additional internship capacity those graduates will need, so "
           "the top and bottom of the pipeline expand together.")

    doc.add_heading("Recommendation 3:  Ring-fence and make house-officer funding portable  "
                    "(targets Cause 3)", level=2)
    bullet("Create a dedicated, ring-fenced federal funding line for house-officer stipends that "
           "**follows the graduate to any accredited centre, public or private**, rather than "
           "being tied to specific federal institutions. Funding should enable training wherever "
           "capacity exists, not confine it to where the payroll happens to sit.")
    bullet("Guarantee timely payment under the existing salary structure, and remove the "
           "budget-capture barrier that currently lets willing hospitals decline interns for "
           "want of a line item.")

    doc.add_heading("Recommendation 4:  Make allocation transparent, rules-based, and clean  "
                    "(targets Cause 4)", level=2)
    bullet("Publish accredited centres, their quotas, and live vacancies in real time on the "
           "portal, so every graduate can see what is available.")
    bullet("Run allocation on transparent, rules-based, automated matching, with a published "
           "appeals and grievance channel.")
    bullet("Pair this with a clear anti-corruption safeguard and reporting route, so any demand "
           "for payment to secure a placement can be reported and acted on. Transparency is the "
           "cheapest and fastest cure for the bribery the current scarcity invites.")

    # ---------------- 5 ----------------
    doc.add_heading("5.  A proven direction, and an open door", level=1)
    para("None of this is speculative. The Nursing and Midwifery Council of Nigeria, facing the "
         "same shortage of internship sites for its own graduates, has moved to accredit a wider "
         "range of hospitals, including in the private sector, as clinical training sites. The "
         "medical profession can follow the same logic.")
    para("More importantly, the MDCN has already told the National Assembly that the solution is "
         "to include state and privately owned hospitals in the housemanship system. The "
         "regulator, the profession, and the private sector are in agreement on the direction. "
         "What is needed now is the political sponsorship to turn that agreement into "
         "accreditation reform, a funding mechanism, and a transparent portal, on a defined "
         "timeline.")

    # ---------------- 6 ----------------
    doc.add_heading("6.  The ask, and the Federation's offer", level=1)
    para("The Healthcare Federation of Nigeria respectfully asks the Honourable Minister of "
         "State to:")
    for i, item in enumerate([
        "**Sponsor a review of housemanship accreditation** to introduce proportionate, "
        "competency-based criteria and a rotational, consortium model, and to expand private and "
        "state hospital participation and quotas.",
        "**Commission a national housemanship curriculum and end-of-internship competency "
        "assessment** to guarantee quality as capacity broadens.",
        "**Direct the creation of a ring-fenced, portable house-officer funding line** so that "
        "funding follows the graduate to any accredited centre.",
        "**Mandate a transparent, real-time, rules-based allocation portal** with "
        "anti-corruption safeguards.",
        "**Convene a time-bound implementation task team** drawing together the Ministry, MDCN, "
        "and the private sector, with quarterly capacity targets reconciled against graduate "
        "output.",
    ], start=1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        add_runs(p, item)
    para("The Federation offers more than a request. Its members operate a substantial share of "
         "Nigeria's accreditable private hospital capacity, and are willing to help design the "
         "rotational model, pilot it, and open their institutions as training sites. The "
         "Federation can also help mobilise senior supervisory capacity, including from Nigerian "
         "medical networks at home and in the diaspora, to meet the consultant-supervision "
         "standards that quality training requires. We are ready to sit down with the relevant "
         "offices and work through the detail carefully.")

    hrule()
    para("Prepared for the Healthcare Federation of Nigeria by Dr Debo Odulana, a member of the "
         "Federation who also serves as President of the Doctors Foundation for Care. Submitted "
         "for the review and adoption of HFN leadership ahead of transmission.",
         italic=True, size=10, color=MUTED)

    # ---------------- ANNEX ----------------
    doc.add_page_break()
    doc.add_heading("Annex:  figures and sources for verification", level=1)
    para("The following figures are drawn from public reporting and official statements, and are "
         "provided so that HFN leadership can verify each before transmission.")
    for b in [
        "6,000 graduates per year; 4,000 system capacity; about 2,000 stranded annually: MDCN "
        "Registrar, Dr Fatima Kyari, presentation to the Senate Committee on Health, reported "
        "February 2026 (ThisDay, New Dawn, Daily Trust).",
        "114 accredited housemanship hospitals (47 teaching, 21 Federal Medical Centres, 37 "
        "general and specialist that are mostly state-owned and military, 9 private); only 44, "
        "all federal, on the centralised portal, the other 70 outside it; about 2,783 medical "
        "and 211 dental places advertised across the 44 portal centres: MDCN Registrar, Dr "
        "Fatima Kyari, quoted in Punch and Punch Healthwise reporting. Directly attributable to "
        "MDCN.",
        "Reports of graduates asked for sums such as N400,000 to secure a place, sought-after "
        "centres with about 30 slots against 300-plus applicants, and teaching hospitals routing "
        "only about 40 per cent of places to the open portal: named-graduate accounts in Punch "
        "and Punch Healthwise reporting.",
        "Core postings of 12 weeks each in Medicine, Surgery, Obstetrics and Gynaecology, and "
        "Paediatrics: MDCN housemanship guidance.",
        "Medical school intake up about 38 per cent since 2023, new schools added, about 20 in "
        "the accreditation pipeline; new intakes up about 160 per cent 2023 to end 2025: Federal "
        "Ministry of Health statements, 2024 to 2025.",
        "More than 8,500 Nigerian doctors joined the UK register 2021 to 2024; an estimated "
        "16,000 doctors lost to emigration over five years against roughly 55,000 licensed; "
        "about half of Nigerian-trained doctors practise abroad: General Medical Council data, "
        "Punch reporting, and related analyses.",
        "Cost of training a doctor in Nigeria estimated in the tens of thousands of dollars: Mo "
        "Ibrahim Foundation analysis.",
        "National Policy on Health Workforce Migration approved by the Federal Executive Council, "
        "August 2024; Nigeria Health Sector Renewal Investment Initiative: Federal Ministry of "
        "Health.",
        "Nursing and Midwifery Council of Nigeria broadening accredited clinical training sites: "
        "NMCN and WHO Africa reporting.",
    ]:
        bullet(b)

    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
